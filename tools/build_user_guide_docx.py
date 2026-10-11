from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "箱门鲜知_用户使用说明书_V1.0.docx"

FONT = "PingFang SC"
GREEN = RGBColor(21, 60, 53)
GREEN_2 = RGBColor(39, 117, 87)
PALE_GREEN = "EAF4EF"
YELLOW = RGBColor(230, 166, 29)
PALE_YELLOW = "FFF6DC"
INK = RGBColor(31, 42, 38)
MUTED = RGBColor(96, 108, 102)
LIGHT_BORDER = "D8E1DC"
WHITE = RGBColor(255, 255, 255)
CONTENT_DXA = 9360


SCREEN = DOCS / "assets" / "screen_preview_f53_portrait.png"
BUTTONS = ROOT / "hardware" / "enclosure_v0_7" / "previews" / "rear_button_labels_outside_v0_7.png"
LAYOUT = ROOT / "hardware" / "enclosure_v0_7" / "previews" / "electronics_layout_v0_7.png"
EXPLODED = ROOT / "hardware" / "enclosure_v0_7" / "previews" / "exploded_view_v0_7.png"
INSTALLED = ROOT / "hardware" / "enclosure_v0_7" / "previews" / "installed_effect_v0_7.png"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Iterable[int], indent: int = 120) -> None:
    widths = list(widths)
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    """Keep a table row intact instead of leaving a blank continuation on the next page."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_run(run, *, size=None, bold=None, color=None, italic=None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_text(doc, text: str, *, style=None, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run(first, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run(rest)
    else:
        run = p.add_run(text)
        set_run(run)
    return p


def add_bullets(doc, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run(p.add_run(item))


def new_decimal_list_id(doc: Document) -> int:
    """Create an independent one-level numbered list that starts at 1."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.append(level_text)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)
    p_pr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "300")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_steps(doc, items: Iterable[str]) -> None:
    num_id = new_decimal_list_id(doc)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id_node)
        p_pr.append(num_pr)
        set_run(p.add_run(item))


def keep_table_together(table) -> None:
    """Keep a short reference table and its rows on one page when possible."""
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if row_index == len(table.rows) - 1:
                    continue
                p_pr = paragraph._p.get_or_add_pPr()
                if p_pr.find(qn("w:keepNext")) is None:
                    p_pr.append(OxmlElement("w:keepNext"))


def add_callout(doc, title: str, body: str, *, caution=False) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_YELLOW if caution else PALE_GREEN)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(title), bold=True, color=YELLOW if caution else GREEN_2)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    set_run(p2.add_run(body), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc, path: Path, caption: str, *, width: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    # Make every instructional figure understandable to screen-reader users.
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    set_run(cap.add_run(caption), size=9, color=MUTED)


def add_heading(doc, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r = OxmlElement("w:r")
    r.append(fld_char1)
    r.append(instr_text)
    r.append(fld_char2)
    paragraph._p.append(r)
    set_run(paragraph.add_run(" 页"), size=9, color=MUTED)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Title": (30, GREEN, 0, 8),
        "Subtitle": (15, MUTED, 0, 24),
        "Heading 1": (16, GREEN, 18, 10),
        "Heading 2": (13, GREEN_2, 14, 7),
        "Heading 3": (12, GREEN_2, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(header.add_run("箱门鲜知  |  用户使用说明书"), size=9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(kicker.add_run("冰箱食材管理墨水屏"), size=11, bold=True, color=YELLOW)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(title.add_run("箱门鲜知"), size=30, bold=True, color=GREEN)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(subtitle.add_run("用户使用说明书"), size=15, color=MUTED)
    add_figure(doc, INSTALLED, "V0.7 设备正面与背面外观示意", width=6.25)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(meta.add_run("版本 V1.0  |  适用设备 V0.7  |  2026-09-02"), size=10, color=MUTED)
    doc.add_page_break()


def add_quick_start(doc: Document) -> None:
    add_heading(doc, "快速开始", 1)
    add_callout(
        doc,
        "第一次使用只需五步",
        "接通电源 → 微信扫码 → 输入 2.4GHz Wi-Fi → 允许蓝牙并开始绑定 → 等待墨水屏刷新。",
    )
    add_steps(
        doc,
        [
            "将设备稳定吸附在冰箱门外侧，并接通电源。",
            "微信扫描设备标签或墨水屏小程序码。",
            "输入 2.4GHz Wi-Fi 名称和密码。",
            "允许蓝牙权限，点击“开始绑定”，保持手机靠近设备。",
            "看到“绑定成功”后等待首次画面刷新，再进入“今日”页核对数据。",
        ],
    )
    add_callout(doc, "请牢记", "RESET 只负责重启；BOOT 用于清除网络配置。正常重启时不要按住 BOOT。", caution=True)


def add_overview(doc: Document) -> None:
    add_heading(doc, "1. 产品用途", 1)
    add_text(doc, "箱门鲜知安装在冰箱门外侧。用户在微信小程序中维护食材、保鲜日期和待办事项；设备每 5 分钟检查一次新画面，并在内容变化时刷新四色墨水屏。")
    add_bullets(doc, ["最多显示 15 项食材及建议食用状态。", "最多显示 3 条待办事项。", "显示当前家庭名称、设备小程序码和画面生成时间。"])
    add_callout(doc, "墨水屏特性", "断电后仍保留最后一幅画面是正常现象，不代表设备仍在联网。")
    add_figure(doc, SCREEN, "墨水屏显示内容示意", width=3.55)


def add_safety(doc: Document) -> None:
    add_heading(doc, "2. 安全须知", 1)
    add_steps(doc, [
        "设备只能安装在冰箱外侧或其他干燥金属表面，不得放入冷藏室、冷冻室或靠近水槽。",
        "只使用合格的 5V USB 电源和完好充电线。",
        "电池出现鼓包、破损、异味或异常发热时，立即停止充电和使用。",
        "不要挤压、刺穿、弯折或自行改接电池；电池接头极性不得接反。",
        "背面含强磁铁，应远离磁卡、机械手表及对磁场敏感的物品。",
        "不要按压屏幕中央或折死 FPC；操作按键孔时只使用圆头塑料杆。",
        "儿童使用时应由成人协助。",
    ])


def add_device_parts(doc: Document) -> None:
    add_heading(doc, "3. 认识设备", 1)
    add_bullets(doc, [
        "正面：7.5 英寸四色墨水屏。",
        "顶部中央：USB-C 充电口。",
        "背面观察窗：充电状态和四格电量指示。",
        "BOOT 孔：配网/清除网络配置。",
        "RESET 孔：重新启动设备。",
        "设备标签：设备编号和小程序码。",
    ])
    add_callout(doc, "方向", "从设备背面看，USB-C 充电口朝上时，BOOT 在左、RESET 在右。")
    add_figure(doc, BUTTONS, "图 1  背面 BOOT 与 RESET 位置", width=6.3)


def add_pairing(doc: Document) -> None:
    add_heading(doc, "4. 首次绑定", 1)
    add_heading(doc, "4.1 准备条件", 2)
    add_bullets(doc, ["已登录微信的手机。", "可以正常上网的 2.4GHz Wi-Fi。", "Wi-Fi 名称和密码。", "设备持续通电，手机距离设备不超过 1 米。"])
    add_heading(doc, "4.2 绑定步骤", 2)
    add_steps(doc, [
        "接通设备电源。未保存配置的设备会自动进入蓝牙配网模式。",
        "微信扫描设备标签或墨水屏上的小程序码。",
        "在“绑定墨水屏”页面输入家庭 2.4GHz Wi-Fi 名称和密码。",
        "允许微信使用蓝牙；Android 手机还可能要求附近设备或定位权限。",
        "点击“开始绑定”，保持手机靠近设备，不切换页面、不关闭电源。",
        "小程序显示“绑定成功”后，等待设备联网并刷新第一幅库存画面。",
        "进入“今日”页，确认家庭名称和墨水屏画面一致。",
    ])
    add_callout(doc, "隐私说明", "Wi-Fi 密码通过加密连接传给设备，不会写入普通业务日志。")


def add_miniapp(doc: Document) -> None:
    add_heading(doc, "5. 小程序日常使用", 1)
    add_heading(doc, "5.1 四个主要页面", 2)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "页面"
    headers[1].text = "主要用途"
    for cell in headers:
        set_cell_shading(cell, PALE_GREEN)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run(run, bold=True, color=GREEN)
    rows = [
        ("今日", "临期、今日到期、已过期食材和家庭摘要。"),
        ("食材", "新增、编辑、消耗、丢弃、删除食材和查看操作历史。"),
        ("待办", "维护墨水屏待办；家庭管理员和普通成员均可编辑。"),
        ("我的", "管理家庭、成员、冰箱名称、设备、提醒和隐私数据。"),
    ]
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = left, right
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                set_run(run)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [1800, 7560])
    keep_table_together(table)
    doc.add_paragraph()

    sections = [
        ("5.2 新增食材", ["进入“食材”，点击“新增食材”。", "填写名称、数量、单位、储存位置、入库日期和建议食用日期。", "点击“保存并更新库存”，返回列表确认。"]),
        ("5.3 快速记一笔", ["输入自然语言，例如“买了两盒牛奶”“牛奶喝完了”或“冷藏室有哪些快过期的东西”。", "写操作必须先检查确认卡片，再点击确认。"]),
        ("5.4 导入订单图片", ["选择一张完整订单图或最多 9 张连续截图。", "开始识别并逐项核对确认卡片。", "修改错误内容后再批量入库。"]),
        ("5.5 设置水墨屏待办", ["进入与“今日、食材、我的”同级的“待办”页。", "最多保存 10 条，每条最多 40 个字符。", "管理员和普通成员均可编辑；屏幕只显示前 3 条。"]),
        ("5.6 修改冰箱名称", ["进入“我的 → 家庭成员”。", "管理员修改名称并保存，最长 40 个字符。", "设备下一次获取新画面后显示新名称。"]),
    ]
    for title, items in sections:
        add_heading(doc, title, 2)
        add_steps(doc, items)


def add_refresh_power(doc: Document) -> None:
    add_heading(doc, "6. 画面更新规则", 1)
    add_bullets(doc, [
        "正常情况下每 5 分钟检查一次新画面。",
        "食材、待办或名称变化时才刷新；没有变化时屏幕不闪烁。",
        "保存后先等待 5 分钟，网络较慢时可再等待一个周期。",
        "最多显示 15 项食材和 3 条待办，其余内容仍保存在小程序。",
        "右上角时间是画面生成的北京时间，不是持续运行的电子时钟。",
    ])
    add_heading(doc, "7. 充电与电源", 1)
    add_text(doc, "设备使用 3.7V 5000mAh 带保护板聚合物锂电池和 Lipo Rider Plus。外部 USB-C 口用于 5V 充电，模块通过 USB-A 向 ESP32 输出 5V。")
    add_steps(doc, [
        "确认设备稳定，充电线不会拉扯外壳。",
        "将 USB-C 公头完全插入顶部中央充电口。",
        "连接合格的 5V USB 电源，建议使用 5V/2A。",
        "通过背面观察窗查看充电状态灯和四格电量灯。",
        "充电时可保持设备工作；出现明显发热时立即断电检查。",
    ])
    add_callout(doc, "电池安全", "不得使用鼓包、破损或没有保护板的电池。普通用户不得拆开外壳更换电池。", caution=True)


def add_reset_transfer(doc: Document) -> None:
    add_heading(doc, "8. 重启、重新配网和解绑", 1)
    add_heading(doc, "8.1 日常重启", 2)
    add_text(doc, "使用圆头塑料杆短按一次 RESET。该操作不会删除 Wi-Fi、绑定关系或屏幕内容。正常重启时不要按住 BOOT。")
    add_heading(doc, "8.2 清除网络并重新配网", 2)
    add_steps(doc, [
        "完全松开 BOOT，短按一次 RESET。",
        "设备启动后会提供约 8 秒的清配置窗口。",
        "在窗口内单独按住 BOOT 约 2 秒，然后松开。",
        "设备清除网络配置并进入蓝牙配网模式。",
        "回到“我的 → 墨水屏设备”，重新执行首次绑定。",
    ])
    add_callout(doc, "禁止操作", "不要按住 BOOT 再按 RESET，否则设备会进入固件下载模式。", caution=True)
    add_heading(doc, "8.3 误入下载模式", 2)
    add_steps(doc, ["完全松开 BOOT。", "只短按一次 RESET。", "仍未恢复时断电，确认 BOOT 没有卡住，等待 5 秒后重新通电。"])
    add_heading(doc, "8.4 解绑或转让设备", 2)
    add_steps(doc, [
        "管理员在“我的 → 墨水屏设备”点击“解绑”。",
        "按 8.2 的步骤清除设备本地配置。",
        "确认设备重新进入蓝牙配网模式。",
        "转让前清除原家庭画面；无法处理时联系设备提供者刷新中性欢迎画面。",
    ])
    add_callout(doc, "注意", "只在小程序解绑而不清除设备配置，不等于完成转让前的数据清理。")


def add_fridge_install(doc: Document) -> None:
    add_heading(doc, "9. 安装到冰箱门", 1)
    add_steps(doc, [
        "擦净冰箱门外侧，去除水分、油污和灰尘。",
        "双手托住设备两侧，使背面平行靠近目标位置。",
        "完整贴合后再松手，不要让一个角先猛烈吸附。",
        "轻推检查是否下滑；有滑动时立即取下检查。",
        "避开门铰链、门缝、冰箱显示面板和高温散热区域。",
        "让充电线自然下垂，不得成为拉动设备的受力点。",
    ])
    add_callout(doc, "防跌落检查", "设备总重约 285g。正式使用前必须在目标冰箱门上连续观察不少于 24 小时，确认不会向下滑动。", caution=True)


def add_assembly(doc: Document) -> None:
    add_heading(doc, "10. V0.7 DIY 装配说明", 1)
    add_text(doc, "已购买整机的用户无需执行本章。V0.7 外形为 120×184×19mm，由前框、中托和后盖组成。任何部件需要拉开、弯曲或靠螺丝强行闭合时，应停止装配。")
    add_figure(doc, EXPLODED, "图 2  前框、中托与后盖爆炸结构", width=6.25)
    add_heading(doc, "10.1 器件位置", 2)
    add_text(doc, "从打开的设备背面看，USB-C 充电口朝上：ESP32 位于左上，Lipo Rider 位于顶部中央，DESPI-C02 位于右侧中部，806090 电池横放在下方。")
    add_figure(doc, LAYOUT, "图 3  中托器件与 9 颗磁条位置", width=4.5)
    add_heading(doc, "10.2 装配顺序", 2)
    add_steps(doc, [
        "验证前框、中托、后盖空壳能徒手自然组合和分离。",
        "将 9 颗 38×9×1.5mm 磁条干装后用极薄环氧胶固定，朝冰箱门的一面保持同一极性。",
        "安装水墨屏，显示面朝窗口，FPC 从右侧开放位置自然伸出。",
        "用 4 颗 M2×4mm 自攻螺丝固定 ESP32，天线朝顶部，BOOT、RESET 与后盖孔对应。",
        "用 4 颗 M2×4mm 自攻螺丝固定 Lipo Rider，USB-C 输入朝顶部，USB-A 输出朝下。",
        "用不超过 0.2mm 的薄双面胶或薄型可拆魔术贴固定 DESPI-C02。",
        "用不超过 0.2mm 的薄双面胶固定电池底面，必要时仅在四周加窄 EVA 条。",
        "核对电池红正黑负，并按接线表连接 ESP32 与 DESPI-C02。",
        "不盖后盖完成电源、Wi-Fi、画面下载和四色刷新测试。",
        "三层自然落位后，用 4 颗 M2×14mm 自攻螺丝按对角顺序分次拧紧。",
    ])
    add_callout(doc, "紧固件限制", "不得使用 M2×16mm 夹紧外壳，也不得用螺丝把没有自然落位的部件强行拉合。", caution=True)


def add_troubleshooting(doc: Document) -> None:
    add_heading(doc, "11. 常见问题", 1)
    rows = [
        ("搜索不到设备", "确认已进入配网模式；开启蓝牙/附近设备权限；手机靠近至 1 米内；重开小程序。"),
        ("Wi-Fi 连接失败", "只使用 2.4GHz；核对名称和密码；靠近路由器；确认双频路由器允许 2.4GHz 设备。"),
        ("保存后屏幕不更新", "等待 5～10 分钟；确认家庭、电量和网络；短按 RESET，禁止同时按 BOOT。"),
        ("断电后仍有画面", "属于墨水屏正常特性；以小程序状态和右上角更新时间判断是否在线。"),
        ("设备向下滑", "立即取下；清洁表面；检查 9 颗磁条和后盖；未解决前禁止悬挂。"),
        ("充电灯不亮/无法开机", "更换确认正常的 5V 电源和线；插到底；仍无反应时联系提供者，不自行拆电池。"),
        ("显示 waiting for download", "松开 BOOT，只短按 RESET；必要时断电 5 秒后重新通电。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.rows[0].cells[0].text = "现象"
    table.rows[0].cells[1].text = "处理方法"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, PALE_GREEN)
        for run in cell.paragraphs[0].runs:
            set_run(run, bold=True, color=GREEN)
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = left, right
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                set_run(run, size=10.2)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [2400, 6960])
    doc.add_paragraph()


def add_care_specs(doc: Document) -> None:
    add_heading(doc, "12. 清洁、存放和运输", 1)
    add_bullets(doc, [
        "使用干燥柔软的布擦拭，不得用水冲洗或使用腐蚀性清洁剂。",
        "搬动时双手托住设备，禁止只提充电线或 FPC。",
        "长期存放应阴凉、干燥，并远离金属碎屑和易燃物。",
        "运输前关闭供电并保护屏幕；锂电池运输遵守承运方规定。",
    ])
    add_heading(doc, "13. 主要规格", 1)
    specs = [
        ("产品外形", "120×184×19mm（V0.7）"),
        ("显示屏", "GDEM075F53；7.5 英寸；800×480；黑白黄红四色"),
        ("主控", "微雪 ESP32-S3-DEV-KIT-N16R8"),
        ("网络", "2.4GHz Wi-Fi；蓝牙配网"),
        ("更新周期", "正常情况下每 5 分钟检查一次"),
        ("电池", "806090；3.7V 5000mAh；带保护板"),
        ("电源模块", "Seeed Lipo Rider Plus 106990290"),
        ("充电输入", "USB-C；5V；建议 5V/2A 电源"),
        ("显示容量", "最多 15 项食材、3 条待办"),
        ("待办容量", "最多 10 条；每条最多 40 个字符"),
        ("固定方式", "9 颗 38×9×1.5mm 钕铁硼硬磁条"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "规格"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, PALE_GREEN)
        for run in cell.paragraphs[0].runs:
            set_run(run, bold=True, color=GREEN)
    for key, value in specs:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, value
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                set_run(run, size=10.2)
    for row in table.rows:
        prevent_row_split(row)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [2400, 6960])


def add_help(doc: Document) -> None:
    add_heading(doc, "14. 获取帮助", 1)
    add_text(doc, "联系设备提供者时，请准备以下信息：")
    add_bullets(doc, ["设备编号。", "小程序完整错误文字。", "墨水屏右上角更新时间。", "问题发生时间。", "是否正在充电及指示灯状态。", "必要时提供串口日志。"])
    add_callout(doc, "不要发送", "Wi-Fi 密码、设备密钥、微信 AppSecret 或其他账号密钥。", caution=True)
    add_heading(doc, "15. 参考资料", 1)
    add_bullets(doc, [
        "Seeed Studio Lipo Rider Plus 官方资料：https://wiki.seeedstudio.com/cn/Lipo-Rider-Plus/",
        "V0.7 装配基线：hardware/enclosure_v0_7/ASSEMBLY_GUIDE.md",
        "设备出厂流程：docs/FACTORY_PROVISIONING.md",
    ])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    set_run(p.add_run("— 说明书结束 —"), size=10, color=MUTED)


def build() -> None:
    doc = Document()
    style_document(doc)
    core = doc.core_properties
    core.title = "箱门鲜知用户使用说明书"
    core.subject = "冰箱食材管理墨水屏设备使用、配网、充电、安装与故障处理"
    core.author = "箱门鲜知"
    core.keywords = "箱门鲜知, 墨水屏, 冰箱食材管理, 使用说明书"
    core.comments = "V1.0"

    add_cover(doc)
    add_quick_start(doc)
    add_overview(doc)
    add_safety(doc)
    add_device_parts(doc)
    add_pairing(doc)
    add_miniapp(doc)
    add_refresh_power(doc)
    add_reset_transfer(doc)
    add_fridge_install(doc)
    add_assembly(doc)
    add_troubleshooting(doc)
    add_care_specs(doc)
    add_help(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
