from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172B3A"
MUTED = "66717A"
LIGHT = "F2F4F7"
CALLOUT = "EAF2F8"
RED = "9B1C1C"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color="D9DEE3", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    """Keep a table row intact so continuation pages retain complete records."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_keep_with_next(paragraph, value=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if value and node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    elif not value and node is not None:
        p_pr.remove(node)


def set_font(run, latin="Arial", east_asia="Arial Unicode MS", size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def apply_inline_markdown(paragraph, text: str, size=11, color=INK, bold=False) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|<https?://[^>]+>)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        is_code = part.startswith("`") and part.endswith("`")
        is_url = part.startswith("<http") and part.endswith(">")
        content = part[2:-2] if is_bold else part[1:-1] if is_code or is_url else part
        run = paragraph.add_run(content)
        set_font(
            run,
            latin="Menlo" if is_code else "Arial",
            east_asia="Arial Unicode MS",
            size=9.5 if is_code else size,
            color=BLUE if is_url else color,
            bold=bold or is_bold,
        )


def paragraph_border_bottom(paragraph, color=BLUE, size="10", space="6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=9, color=MUTED)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in specs.items():
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def configure_section(doc: Document, meta: dict[str, str]) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    version = meta.get("version", "V1.0")
    r = p.add_run(f"冰箱食材管理设备  ·  PRD {version}")
    set_font(r, size=9, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run("第 ")
    set_font(r, size=9, color=MUTED)
    add_page_number(fp)
    r = fp.add_run(" 页")
    set_font(r, size=9, color=MUTED)


def parse_frontmatter(lines: list[str]) -> tuple[dict[str, str], list[str]]:
    meta: dict[str, str] = {}
    if not lines or lines[0].strip() != "---":
        return meta, lines
    idx = 1
    while idx < len(lines) and lines[idx].strip() != "---":
        if ":" in lines[idx]:
            key, value = lines[idx].split(":", 1)
            meta[key.strip()] = value.strip()
        idx += 1
    return meta, lines[idx + 1 :]


def add_masthead(doc: Document, meta: dict[str, str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(meta.get("title", "冰箱食材管理设备"))
    set_font(r, size=24, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(meta.get("subtitle", "第一版产品需求文档（MVP）"))
    set_font(r, size=14, color=MUTED, bold=True)

    rows = [
        ("版本", meta.get("version", "V1.0")),
        ("日期", meta.get("date", "2026-07-21")),
        ("状态", meta.get("status", "需求基线")),
        ("产品阶段", "个人 DIY 原型 / 后续开源"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{label}：")
        set_font(r, size=10.5, color=INK, bold=True)
        r = p.add_run(value)
        set_font(r, size=10.5, color=INK)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(10)
    rule.paragraph_format.space_after = Pt(12)
    paragraph_border_bottom(rule, color=BLUE, size="14", space="5")


def table_widths(headers: list[str]) -> list[int]:
    count = len(headers)
    if count == 2:
        return [2200, 7160]
    if count == 3:
        if headers[0] in {"编号", "状态", "版本", "阶段", "对象", "项目", "风险", "模块"}:
            return [1500, 3000, 4860]
        return [2200, 3000, 4160]
    return [CONTENT_WIDTH_DXA // count] * count


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    set_table_geometry(table, table_widths(rows[0]))
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])

    for i, values in enumerate(rows):
        set_row_cant_split(table.rows[i])
        for j, value in enumerate(values):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if j == 0 and len(value) <= 16:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            apply_inline_markdown(p, value, size=9.3, color=INK, bold=(i == 0))
            if i == 0:
                set_cell_shading(cell, LIGHT)

    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)


def markdown_to_docx(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    meta, lines = parse_frontmatter(lines)

    doc = Document()
    configure_styles(doc)
    configure_section(doc, meta)
    add_masthead(doc, meta)

    idx = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(x.strip() for x in paragraph_buffer).strip()
        if text:
            p = doc.add_paragraph()
            apply_inline_markdown(p, text)
        paragraph_buffer = []

    while idx < len(lines):
        line = lines[idx].rstrip()
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            idx += 1
            continue

        if stripped.startswith("#"):
            flush_paragraph()
            match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
            if match:
                level = len(match.group(1))
                heading_text = match.group(2)
                # The acceptance matrix is intentionally paginated in bounded
                # chunks below. Start the section on a clean page so no row is
                # pushed into the header/footer area by office renderers.
                if level == 1 and heading_text.startswith("14. 验收标准"):
                    doc.add_page_break()
                p = doc.add_paragraph(style=f"Heading {level}")
                apply_inline_markdown(p, heading_text, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)
                set_keep_with_next(p)
            idx += 1
            continue

        if stripped.startswith("> "):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), CALLOUT)
            p_pr.append(shd)
            apply_inline_markdown(p, stripped[2:], size=10.5, color=DARK_BLUE, bold=True)
            idx += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            table_lines = []
            while idx < len(lines) and lines[idx].strip().startswith("|"):
                table_lines.append(lines[idx].strip())
                idx += 1
            parsed = []
            for tline in table_lines:
                cells = [c.strip() for c in tline.strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", c or "---") for c in cells):
                    continue
                parsed.append(cells)
            if parsed and parsed[0] == ["编号", "验收项", "通过条件"] and len(parsed) > 9:
                for chunk_start in range(1, len(parsed), 8):
                    if chunk_start > 1:
                        doc.add_page_break()
                        continuation = doc.add_paragraph(style="Heading 2")
                        apply_inline_markdown(
                            continuation,
                            "14. 验收标准（续）",
                            size=13,
                            color=BLUE,
                            bold=True,
                        )
                        set_keep_with_next(continuation)
                    add_table(doc, [parsed[0], *parsed[chunk_start : chunk_start + 8]])
            else:
                add_table(doc, parsed)
            continue

        if re.match(r"^-\s+", stripped):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            apply_inline_markdown(p, re.sub(r"^-\s+", "", stripped))
            idx += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            flush_paragraph()
            p = doc.add_paragraph(style="List Number")
            apply_inline_markdown(p, re.sub(r"^\d+\.\s+", "", stripped))
            idx += 1
            continue

        paragraph_buffer.append(stripped)
        idx += 1

    flush_paragraph()

    core = doc.core_properties
    core.title = meta.get("title", "冰箱食材管理设备") + " " + meta.get("subtitle", "PRD")
    core.subject = "第一版产品需求文档"
    core.author = "项目团队"
    core.keywords = "冰箱, 食材管理, 微信小程序, 电子墨水屏, ESP32, PRD"
    core.comments = f"根据 {meta.get('date', '2026-07-21')} 需求基线整理"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_prd_docx.py <input.md> <output.docx>")
    markdown_to_docx(Path(sys.argv[1]), Path(sys.argv[2]))
